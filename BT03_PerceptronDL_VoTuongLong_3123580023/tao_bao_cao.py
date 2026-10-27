from datetime import date
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

import p1_cau1_ranh_gioi as p1c1
import p1_cau2_perceptron_hinh_ve as p1c2
import p1_cau3_hoc_ham_or as p1c3
import p1_cau4_hoc_ham_and as p1c4
import p2_cau1_2_mien_kich_hoat as p2c12
import p2_cau3_4_xet_danh_hieu as p2c34
from ve_hinh import bieu_thuc_tuyen_tinh as btt
from ve_hinh import so

HO_TEN = "Võ Tường Long"
GIANG_VIEN = "Đỗ Như Tài"
MA_SV = "3123580023"
TEN_FILE = "BT03_PerceptronDL_VoTuongLong_3123580023.docx"
NHANH = "BT03_PerceptronDL"
REPO = f"https://github.com/tuonglong1012/Mon_Hoc_Sau_DDU1231/tree/{NHANH}"
THU_MUC = Path(__file__).resolve().parent
MAU_TIEU_DE = RGBColor(0x00, 0x00, 0x00)
LOGO = THU_MUC / "hinh" / "logoSGU.jpg"
DEM_HINH = [0]


def muc_luc(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("MỤC LỤC")
    r.bold = True
    r.font.size = Pt(16)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(1)
    run = p.add_run()
    for loai, noi_dung in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None), ("text", "Nhấn chuột phải vào đây và chọn Update Field để cập nhật mục lục."), ("end", None)):
        if loai == "text":
            el = OxmlElement("w:t")
            el.text = noi_dung
        elif loai:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), loai)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = noi_dung
        run._r.append(el)


def cai_dat_tai_lieu(doc):
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
    kieu = doc.styles["Normal"]
    kieu.font.name = "Times New Roman"
    kieu.font.size = Pt(13)
    kieu.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    kieu.paragraph_format.space_after = Pt(4)
    kieu.paragraph_format.line_spacing = 1.2
    for ten, co in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 13)):
        h = doc.styles[ten]
        h.font.name = "Times New Roman"
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        h.font.size = Pt(co)
        h.font.bold = True
        h.font.color.rgb = MAU_TIEU_DE
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)


def them_so_trang(doc):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for loai, noi_dung in (("begin", None), (None, "PAGE"), ("end", None)):
        if loai:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), loai)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = noi_dung
        run._r.append(el)


def doan(doc, van_ban, dam=False, nghieng=False, can=None, co=None):
    p = doc.add_paragraph()
    run = p.add_run(van_ban)
    run.bold = dam
    run.italic = nghieng
    if co:
        run.font.size = Pt(co)
    if can:
        p.alignment = can
    return p


def doan_nhieu_phan(doc, cac_phan):
    p = doc.add_paragraph()
    for noi_dung, dam in cac_phan:
        r = p.add_run(noi_dung)
        r.bold = dam
    return p


def cong_thuc(doc, van_ban):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(van_ban)
    run.font.name = "Cambria Math"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria Math")
    run.font.size = Pt(13)
    return p


def gach_dau_dong(doc, van_ban):
    return doc.add_paragraph(van_ban, style="List Bullet")


def danh_so(doc, van_ban):
    return doc.add_paragraph(van_ban, style="List Number")


def bang(doc, tieu_de_cot, cac_dong, do_rong=None, chu_nho=False):
    t = doc.add_table(rows=1, cols=len(tieu_de_cot))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, ten in enumerate(tieu_de_cot):
        o = t.rows[0].cells[i]
        o.text = ""
        r = o.paragraphs[0].add_run(ten)
        r.bold = True
        r.font.size = Pt(10 if chu_nho else 11)
        o.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for dong in cac_dong:
        cells = t.add_row().cells
        for i, gt in enumerate(dong):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(gt))
            r.font.size = Pt(10 if chu_nho else 11)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if do_rong:
        for dong in t.rows:
            for i, w in enumerate(do_rong):
                dong.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def hinh(doc, duong_dan, ten, rong=16):
    DEM_HINH[0] += 1
    doc.add_picture(str(duong_dan), width=Cm(rong))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doan(doc, f"Hình {DEM_HINH[0]}. {ten}", nghieng=True, can=WD_ALIGN_PARAGRAPH.CENTER, co=11)


def khoi_code(doc, ma):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    o = t.rows[0].cells[0]
    o.text = ""
    p = o.paragraphs[0]
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(0)
    dong = ma.rstrip("\n").split("\n")
    for k, d in enumerate(dong):
        r = p.add_run(d)
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        r.font.size = Pt(9)
        if k < len(dong) - 1:
            r.add_break(WD_BREAK.LINE)
    doc.add_paragraph()


def cat_logo():
    with Image.open(LOGO) as anh:
        anh = anh.convert("RGB")
        mat_na = Image.eval(anh.convert("L"), lambda p: 255 if p < 235 else 0)
        khung = mat_na.getbbox()
        anh = anh.crop(khung)
        canh = max(anh.size)
        vuong = Image.new("RGB", (canh, canh), "white")
        vuong.paste(anh, ((canh - anh.width) // 2, (canh - anh.height) // 2))
        duong_dan = THU_MUC / "hinh" / "logoSGU_cat.png"
        vuong.save(duong_dan)
    return duong_dan


def khung_trang(section):
    vien = OxmlElement("w:pgBorders")
    vien.set(qn("w:offsetFrom"), "page")
    for canh in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{canh}")
        el.set(qn("w:val"), "thinThickSmallGap")
        el.set(qn("w:sz"), "24")
        el.set(qn("w:space"), "24")
        el.set(qn("w:color"), "000000")
        vien.append(el)
    section._sectPr.find(qn("w:pgMar")).addnext(vien)


def danh_so_tu_dau(section):
    so_trang = OxmlElement("w:pgNumType")
    so_trang.set(qn("w:start"), "1")
    cols = section._sectPr.find(qn("w:cols"))
    if cols is not None:
        cols.addprevious(so_trang)
    else:
        section._sectPr.append(so_trang)


def dong_bia(doc, van_ban, co, dam=True, nghieng=False, mau=None, truoc=0, sau=0):
    p = doan(doc, van_ban, dam=dam, nghieng=nghieng, can=WD_ALIGN_PARAGRAPH.CENTER, co=co)
    p.paragraph_format.space_before = Pt(truoc)
    p.paragraph_format.space_after = Pt(sau)
    p.paragraph_format.line_spacing = 1.0
    if mau:
        p.runs[0].font.color.rgb = mau
    return p


def trang_bia(doc):
    dong_bia(doc, "TRƯỜNG ĐẠI HỌC SÀI GÒN", 15, sau=2)
    dong_bia(doc, "KHOA TOÁN ỨNG DỤNG", 14)
    dong_bia(doc, "──────── o0o ────────", 12, dam=False, sau=18)
    doc.add_picture(str(cat_logo()), width=Cm(5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    dong_bia(doc, "MÔN HỌC: HỌC SÂU (DEEP LEARNING)", 15, truoc=24, sau=24)
    dong_bia(doc, "BÁO CÁO BÀI TẬP 03", 20, sau=8)
    dong_bia(doc, "PERCEPTRON", 24, mau=MAU_TIEU_DE)
    dong_bia(doc, "VÀ MẠNG NƠ-RON", 24, mau=MAU_TIEU_DE, sau=6)
    dong_bia(doc, "(Perceptron & Neural Networks)", 14, dam=False, nghieng=True, sau=80)
    t = doc.add_table(rows=4, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (nhan, gt) in enumerate((("Giảng viên hướng dẫn:", GIANG_VIEN), ("Sinh viên thực hiện:", HO_TEN), ("Mã số sinh viên:", MA_SV), ("Mã nguồn:", f"github.com/tuonglong1012/Mon_Hoc_Sau_DDU1231\n(nhánh {NHANH})"))):
        t.rows[i].cells[0].text = ""
        r = t.rows[i].cells[0].paragraphs[0].add_run(nhan)
        r.bold = True
        r.font.size = Pt(14)
        t.rows[i].cells[1].text = ""
        r = t.rows[i].cells[1].paragraphs[0].add_run(gt)
        r.font.size = Pt(14 if i < 3 else 12)
        r.bold = i < 2
        t.rows[i].cells[0].width = Cm(6)
        t.rows[i].cells[1].width = Cm(10)
    dong_bia(doc, f"TP. Hồ Chí Minh, tháng {date.today().month} năm {date.today().year}", 13, dam=False, nghieng=True, truoc=110)
    doc.add_section(WD_SECTION.NEW_PAGE)
    bia = doc.sections[0]
    khung_trang(bia)
    bia.different_first_page_header_footer = True
    danh_so_tu_dau(doc.sections[1])


def nhan_pm(y):
    return "+1" if y > 0 else "−1"


def viet_p1_cau1(doc):
    doc.add_heading("PHẦN 1 – CÂU 1. THIẾT KẾ PERCEPTRON CHO RANH GIỚI x₂ = −1.5x₁ + 3", level=1).paragraph_format.page_break_before = True
    doan(doc, "Perceptron đơn hai đầu vào tính tổng s = w₁x₁ + w₂x₂ + b rồi đưa qua hàm bước τ(s) = +1 nếu s ≥ 0, −1 nếu s < 0. Ranh giới quyết định là tập điểm có s = 0, tức đường thẳng w₁x₁ + w₂x₂ + b = 0. Vì vậy chỉ cần đưa phương trình đề cho về dạng này.")

    doc.add_heading("a) Đưa ranh giới về dạng w₁x₁ + w₂x₂ + b = 0", level=2)
    cong_thuc(doc, "x₂ = −1.5x₁ + 3   ⇔   1.5x₁ + x₂ − 3 = 0   ⇔ (nhân 2)   3x₁ + 2x₂ − 6 = 0")
    doan(doc, f"Kiểm tra: mọi điểm (x₁, −1.5x₁ + 3) đều cho 3x₁ + 2(−1.5x₁ + 3) − 6 = 0 (chương trình kiểm tra trên 13 điểm: {p1c1.nam_tren_duong(np.linspace(-2, 4, 13))}). Đường thẳng cắt trục x₂ tại (0, 3) và cắt trục x₁ tại (2, 0).")

    doc.add_heading("b) Perceptron thiết kế", level=2)
    cong_thuc(doc, "w = (w₁, w₂) = (3, 2),    b = −6")
    cong_thuc(doc, "y = τ(3x₁ + 2x₂ − 6) = +1 nếu 3x₁ + 2x₂ − 6 ≥ 0;   −1 nếu 3x₁ + 2x₂ − 6 < 0")
    doan(doc, "Bias b = −6 là trọng số của đầu vào hằng bằng 1. Vector w = (3, 2) là pháp tuyến của ranh giới và chỉ về phía lớp +1 (phía trên – bên phải đường thẳng); gốc tọa độ thuộc lớp −1.")
    hinh(doc, p1c1.ve_hinh(), "Perceptron y = τ(3x₁ + 2x₂ − 6) và ranh giới x₂ = −1.5x₁ + 3")

    doc.add_heading("c) Kiểm tra trên một số điểm", level=2)
    bang(doc, ["Điểm (x₁, x₂)", "s = 3x₁ + 2x₂ − 6", "y = τ(s)", "Vị trí"], [[f"({so(d['x'][0])}, {so(d['x'][1])})", so(d["s"]), nhan_pm(d["y"]), "Trên ranh giới" if d["s"] == 0 else ("Phía lớp +1" if d["s"] > 0 else "Phía lớp −1")] for d in p1c1.bang_kiem_tra()])
    doan(doc, "Nhận xét: nhân cả (w, b) với một số k > 0 bất kỳ, ví dụ (1.5, 1, −3), cho cùng một perceptron. Nhân với k < 0, ví dụ (−3, −2, 6), vẫn cùng ranh giới nhưng đổi vai trò hai lớp. Đề không quy định phía nào là +1 nên chọn (3, 2, −6).")


def viet_p1_cau2(doc):
    doc.add_heading("PHẦN 1 – CÂU 2. PERCEPTRON VÀ ĐƯỜNG THẲNG PHÂN LỚP", level=1).paragraph_format.page_break_before = True
    doan(doc, "Perceptron: y = f_g(wᵀx + b) = +1 nếu wᵀx + b ≥ 0, −1 nếu wᵀx + b < 0. Ranh giới quyết định là đường thẳng wᵀx + b = 0.")

    a, c = p1c2.ranh_gioi_cau_a()
    doc.add_heading("a) Phương trình ranh giới của perceptron trong hình", level=2)
    doan(doc, "Theo hình: trọng số của đầu vào hằng 1 là −3 (bias b = −3), của x₁ là 2.0, của x₂ là 1.0.")
    cong_thuc(doc, "s = 2x₁ + x₂ − 3 = 0   ⇔   x₂ = −2x₁ + 3")
    doan(doc, f"Đường thẳng có hệ số góc {so(a)}, cắt trục x₂ tại (0, {so(c)}) và trục x₁ tại (1.5, 0). Phía trên – bên phải đường thẳng cho y = +1, phía chứa gốc tọa độ cho y = −1.")

    doc.add_heading("b) Đầu ra với x⁽¹⁾ = (−1, 0) và x⁽²⁾ = (2, 0)", level=2)
    s, y = p1c2.dau_ra(p1c2.W_A, p1c2.B_A, p1c2.X_THU)
    bang(doc, ["Đầu vào", "s = 2x₁ + x₂ − 3", "y = f_g(s)"], [[f"x⁽{k + 1}⁾ = ({so(x[0])}, {so(x[1])})", f"2·({so(x[0])}) + {so(x[1])} − 3 = {so(si)}", nhan_pm(yi)] for k, (x, si, yi) in enumerate(zip(p1c2.X_THU, s, y))])
    hinh(doc, p1c2.ve_cau_a(), "Perceptron đề cho, ranh giới x₂ = −2x₁ + 3 và hai điểm x⁽¹⁾, x⁽²⁾")

    w, b = p1c2.perceptron_cau_c()
    doc.add_heading("c) Perceptron ứng với đường thẳng trong hình", level=2)
    doan(doc, "Đường thẳng trong hình đi qua hai điểm (2, 0) trên trục x₁ và (0, −1) trên trục x₂. Hệ số góc (0 − (−1)) / (2 − 0) = 0.5, tung độ gốc −1:")
    cong_thuc(doc, "x₂ = 0.5x₁ − 1   ⇔   x₁ − 2x₂ − 2 = 0")
    doan(doc, f"Đồng nhất với w₁x₁ + w₂x₂ + b = 0 ta được perceptron: trọng số đầu vào hằng 1 là b = {so(b)}, của x₁ là w₁ = {so(w[0])}, của x₂ là w₂ = {so(w[1])}.")
    cong_thuc(doc, f"y = f_g({btt([w[0], w[1], b], ['x₁', 'x₂', ''])})")
    doan(doc, "Với lựa chọn này, phía dưới – bên phải đường thẳng (ví dụ điểm (3, 0): s = 1) cho y = +1, phía chứa gốc tọa độ (s = −2) cho y = −1. Bộ (−1, 2, 2) cũng là một đáp án đúng nếu muốn đảo hai lớp.")
    hinh(doc, p1c2.ve_cau_c(), "Perceptron ứng với đường thẳng đi qua (2, 0) và (0, −1)")

    doc.add_heading("d) Perceptron cho hàm AND và OR", level=2)
    doan(doc, "Mã hóa: đầu vào x₁, x₂ ∈ {0, 1}, đầu ra +1 (đúng) và −1 (sai), bốn mẫu a = (0, 0), b = (1, 0), c = (0, 1), d = (1, 1). Hình thứ nhất trong đề (b, c, d cùng lớp, a khác lớp) là hàm OR; hình thứ hai (chỉ d khác lớp) là hàm AND. Cần chọn đường thẳng tách đúng các mẫu, ở đây chọn đường song song với đường x₁ + x₂ = 0.")
    gach_dau_dong(doc, "AND: chỉ d = (1, 1) cho +1. Đường x₁ + x₂ = 1.5 nằm giữa {a, b, c} (tổng ≤ 1) và d (tổng = 2) → w = (1, 1), b = −1.5.")
    gach_dau_dong(doc, "OR: chỉ a = (0, 0) cho −1. Đường x₁ + x₂ = 0.5 nằm giữa a (tổng = 0) và {b, c, d} (tổng ≥ 1) → w = (1, 1), b = −0.5.")
    cac_dong = []
    for d_and, d_or in zip(p1c2.bang_logic(p1c2.W_AND, p1c2.B_AND, p1c2.T_AND), p1c2.bang_logic(p1c2.W_OR, p1c2.B_OR, p1c2.T_OR)):
        cac_dong.append([d_and["ten"], f"({d_and['x'][0]}, {d_and['x'][1]})", so(d_and["s"]), nhan_pm(d_and["y"]), nhan_pm(d_and["t"]), so(d_or["s"]), nhan_pm(d_or["y"]), nhan_pm(d_or["t"])])
    bang(doc, ["Mẫu", "(x₁, x₂)", "s_AND = x₁ + x₂ − 1.5", "y_AND", "AND", "s_OR = x₁ + x₂ − 0.5", "y_OR", "OR"], cac_dong, chu_nho=True)
    doan(doc, "Cả hai perceptron đều cho đầu ra trùng với bảng chân trị. Hai perceptron chỉ khác nhau ở bias: bias càng âm thì càng cần nhiều đầu vào bằng 1 mới kích hoạt (AND cần cả hai, OR chỉ cần một).")
    hinh(doc, p1c2.ve_cau_d(), "Ranh giới quyết định của perceptron AND (trái) và OR (phải)")


def bang_buoc_or(doc, buoc):
    cac_dong = []
    for i in range(4):
        x = p1c3.X[i]
        cac_dong.append([p1c3.TEN_MAU[i], f"({so(x[0])}, {so(x[1])})", nhan_pm(p1c3.T[i]), so(round(buoc["s"][i], 6)), nhan_pm(buoc["y"][i]), "Sai" if buoc["sai"][i] else "Đúng"])
    bang(doc, ["Mẫu", "x", "t", "s = wᵀx + b", "y = sign(s)", "Kết quả"], cac_dong)


def viet_p1_cau3(doc):
    doc.add_heading("PHẦN 1 – CÂU 3. HỌC HÀM OR BẰNG GRADIENT DESCENT", level=1).paragraph_format.page_break_before = True
    doan(doc, "Tập huấn luyện: a = (0, 0) có t = −1; b = (1, 0), c = (0, 1), d = (1, 1) có t = +1. Perceptron y = sign(wᵀx + b) với sign(s) = +1 nếu s ≥ 0, −1 nếu s < 0. Hàm loss:")
    cong_thuc(doc, "J(w, b) = Σ_{xᵢ ∈ Y} (−tᵢ)(wᵀxᵢ + b),   Y = tập mẫu bị phân loại sai")
    doan(doc, "Với mẫu bị phân loại sai, tᵢ và wᵀxᵢ + b không cùng dấu nên mỗi số hạng −tᵢ(wᵀxᵢ + b) ≥ 0. Do đó J ≥ 0 và J = 0 khi không còn mẫu sai.")

    doc.add_heading("a) Đạo hàm và công thức cập nhật", level=2)
    doan(doc, "Với tập Y cố định, mỗi số hạng là hàm tuyến tính theo w và b: ∂[−tᵢ(wᵀxᵢ + b)]/∂w = −tᵢxᵢ, ∂[−tᵢ(wᵀxᵢ + b)]/∂b = −tᵢ. Cộng trên Y:")
    cong_thuc(doc, "∂J/∂w = −Σ_{xᵢ ∈ Y} tᵢxᵢ,        ∂J/∂b = −Σ_{xᵢ ∈ Y} tᵢ")
    doan(doc, "Gradient descent với tốc độ học ρ đi ngược hướng gradient:")
    cong_thuc(doc, "w(k+1) = w(k) − ρ·∂J/∂w = w(k) + ρ·Σ_{xᵢ ∈ Y} tᵢxᵢ")
    cong_thuc(doc, "b(k+1) = b(k) − ρ·∂J/∂b = b(k) + ρ·Σ_{xᵢ ∈ Y} tᵢ")
    doan(doc, "Ý nghĩa: mẫu lớp +1 bị sai kéo w về phía xᵢ, mẫu lớp −1 bị sai đẩy w ra xa xᵢ.")

    doc.add_heading("b) Thuật toán perceptron learning (batch mode)", level=2)
    danh_so(doc, "Khởi tạo w(0), b(0), chọn tốc độ học ρ > 0, đặt k = 0.")
    danh_so(doc, "Với w(k), b(k): tính sᵢ = w(k)ᵀxᵢ + b(k) và yᵢ = sign(sᵢ) cho toàn bộ N mẫu; lập tập mẫu sai Y = {xᵢ : yᵢ ≠ tᵢ}.")
    danh_so(doc, "Nếu Y rỗng: dừng, perceptron đã phân loại đúng mọi mẫu.")
    danh_so(doc, "Tính gradient trên cả tập Y (một lần cho cả lô): ∂J/∂w = −Σ tᵢxᵢ, ∂J/∂b = −Σ tᵢ.")
    danh_so(doc, "Cập nhật w(k+1) = w(k) + ρΣ tᵢxᵢ, b(k+1) = b(k) + ρΣ tᵢ; tăng k thêm 1 rồi quay lại bước 2 (dừng nếu vượt số vòng lặp tối đa).")
    doan(doc, "Khác với chế độ online (cập nhật ngay sau từng mẫu), batch mode cộng đóng góp của mọi mẫu sai rồi mới cập nhật một lần. Nếu dữ liệu khả tách tuyến tính (như hàm OR), thuật toán dừng sau hữu hạn bước (định lý hội tụ perceptron).")

    lich_su = p1c3.huan_luyen()
    b0 = lich_su[0]
    doc.add_heading("c) Tham số ban đầu ứng với đường thẳng x₂ = 0.5", level=2)
    doan(doc, "Đường ① trong hình là đường nằm ngang cắt trục tung tại 0.5 (đề ghi “y = 0.5”, ở đây trục tung là x₂):")
    cong_thuc(doc, "x₂ = 0.5   ⇔   0·x₁ + 1·x₂ − 0.5 = 0")
    cong_thuc(doc, f"w(0) = (w₁(0), w₂(0))ᵀ = ({so(p1c3.W0[0])}, {so(p1c3.W0[1])})ᵀ,     b(0) = {so(p1c3.B0)}")
    doan(doc, "Chọn dấu sao cho phía trên đường thẳng (chứa c, d có t = +1) cho y = +1. Khi đó a đúng (y = −1), c và d đúng (y = +1), còn b = (1, 0) nằm dưới đường nên bị phân loại sai.")

    doc.add_heading("d) Chạy thuật toán với ρ = 0.4", level=2)
    doan_nhieu_phan(doc, [(f"Bước 0: ", True), (f"w(0) = ({so(b0['w'][0])}, {so(b0['w'][1])}), b(0) = {so(b0['b'])}.", False)])
    bang_buoc_or(doc, b0)
    sai = [p1c3.TEN_MAU[i] for i in np.where(b0["sai"])[0]]
    doan(doc, f"Tập mẫu sai Y = {{{', '.join(sai)}}}. Loss J = −(+1)·({so(b0['s'][1])}) = {so(b0['J'])}.")
    cong_thuc(doc, f"∂J/∂w = −(+1)·(1, 0) = ({so(b0['gw'][0])}, {so(b0['gw'][1])}),      ∂J/∂b = −(+1) = {so(b0['gb'])}")
    cong_thuc(doc, f"w(1) = (0, 1) + 0.4·(1, 0) = ({so(b0['w_moi'][0])}, {so(b0['w_moi'][1])}),      b(1) = −0.5 + 0.4·1 = {so(round(b0['b_moi'], 6))}")
    for buoc in lich_su[1:]:
        doan_nhieu_phan(doc, [(f"Bước {buoc['k']}: ", True), (f"w({buoc['k']}) = ({so(buoc['w'][0])}, {so(round(buoc['w'][1], 6))}), b({buoc['k']}) = {so(round(buoc['b'], 6))}.", False)])
        bang_buoc_or(doc, buoc)
        if buoc["sai"].any():
            doan(doc, f"Còn {int(buoc['sai'].sum())} mẫu sai, J = {so(round(buoc['J'], 6))}; tiếp tục cập nhật.")
        else:
            doan(doc, "Không còn mẫu nào bị phân loại sai (Y rỗng, J = 0) nên thuật toán dừng.")
    cuoi = lich_su[-1]
    doan(doc, f"Kết quả: sau {len(lich_su) - 1} lần cập nhật, perceptron y = sign({btt([cuoi['w'][0], cuoi['w'][1], round(cuoi['b'], 6)], ['x₁', 'x₂', ''])}) phân loại đúng cả 4 mẫu của hàm OR. Đường phân lớp {btt([cuoi['w'][0], cuoi['w'][1], round(cuoi['b'], 6)], ['x₁', 'x₂', ''])} = 0 cắt trục x₁ tại 0.25 và trục x₂ tại 0.1, tách a khỏi b, c, d. Chỉ cần một bước vì ban đầu chỉ có b bị sai; lần cập nhật kéo đường phân lớp nghiêng xuống phía b.")
    hinh(doc, p1c3.ve_cac_buoc(lich_su), "Đường phân lớp tại mỗi bước chạy thuật toán perceptron learning (mẫu sai được khoanh đỏ)", 14)


def viet_p1_cau4(doc):
    cac_dong, w_cuoi_epoch = p1c4.huan_luyen()
    w = w_cuoi_epoch[-1]
    so_epoch = cac_dong[-1]["epoch"]
    doc.add_heading("PHẦN 1 – CÂU 4. HUẤN LUYỆN PERCEPTRON CHO HÀM AND", level=1).paragraph_format.page_break_before = True

    doc.add_heading("a) Thiết kế perceptron", level=2)
    doan(doc, "Perceptron hai đầu vào x₁, x₂ ∈ {0, 1}, trọng số w₁, w₂, ngưỡng θ và hàm bước:")
    cong_thuc(doc, "Y = Step(x₁w₁ + x₂w₂ − θ),   Step(z) = 1 nếu z ≥ 0;  0 nếu z < 0")
    doan(doc, f"Khởi tạo w₁ = {so(p1c4.W0[0])}, w₂ = {so(p1c4.W0[1])}, θ = {so(p1c4.THETA)}, tốc độ học α = {so(p1c4.ALPHA)}. Mỗi mẫu được đưa vào lần lượt (cập nhật online). Sai số e(t) = Y_d − Y, luật học:")
    cong_thuc(doc, "Δwᵢ = α·e(t)·xᵢ,      wᵢ(mới) = wᵢ + Δwᵢ")
    doan(doc, "Theo đề, chỉ w₁, w₂ được cập nhật, ngưỡng θ = 0.4 giữ cố định. Khi e = 0 thì trọng số không đổi; khi e = 1 (dự đoán thiếu) trọng số của đầu vào bằng 1 tăng 0.2; khi e = −1 (dự đoán thừa) giảm 0.2. Một epoch là một lượt qua đủ 4 mẫu; dừng khi trọn một epoch không có lỗi.")

    doc.add_heading("b) Bảng cập nhật trọng số qua từng epoch", level=2)
    doan(doc, "Cột z được tính bằng trọng số trước khi cập nhật; cột w₁, w₂ là trọng số sau khi xử lý mẫu đó.")
    bang(doc, ["Epoch", "x₁", "x₂", "Y_d", "z", "Y", "Sai số e", "w₁", "w₂", "Trạng thái"], [[d["epoch"], d["x1"], d["x2"], d["y_d"], so(round(d["z"], 4)), d["y"], so(d["e"]), so(round(d["w1"], 4)), so(round(d["w2"], 4)), "Cập nhật" if d["cap_nhat"] else "Giữ nguyên"] for d in cac_dong], do_rong=[1.3, 0.9, 0.9, 1.1, 1.4, 0.9, 1.6, 1.3, 1.3, 2.4], chu_nho=True)
    for k in range(1, so_epoch + 1):
        loi = [d for d in cac_dong if d["epoch"] == k and d["cap_nhat"]]
        mo_ta = "không có lỗi → dừng huấn luyện" if not loi else f"{len(loi)} lần cập nhật, cuối epoch w = ({so(w_cuoi_epoch[k][0])}, {so(w_cuoi_epoch[k][1])})"
        gach_dau_dong(doc, f"Epoch {k}: {mo_ta}.")

    doc.add_heading("c) Kết quả và kiểm tra", level=2)
    doan(doc, f"Sau {so_epoch} epoch, perceptron hội tụ với w₁ = {so(w[0])}, w₂ = {so(w[1])}, θ = {so(p1c4.THETA)}:")
    cong_thuc(doc, f"Y = Step({so(w[0])}x₁ + {so(w[1])}x₂ − {so(p1c4.THETA)})")
    kiem_tra = [d for d in cac_dong if d["epoch"] == so_epoch]
    bang(doc, ["x₁", "x₂", f"z = {so(w[0])}x₁ + {so(w[1])}x₂ − {so(p1c4.THETA)}", "Y", "Y_d (AND)"], [[d["x1"], d["x2"], so(round(d["z"], 4)), d["y"], d["y_d"]] for d in kiem_tra])
    doan(doc, f"Cả 4 đầu ra trùng với bảng chân trị AND. Về hình học, ranh giới {so(w[0])}x₁ + {so(w[1])}x₂ = {so(p1c4.THETA)} tách điểm (1, 1) khỏi ba điểm còn lại (w₁ < θ, w₂ < θ nhưng w₁ + w₂ ≥ θ).")
    hinh(doc, p1c4.ve_cac_epoch(w_cuoi_epoch), "Đường ranh giới qua các epoch (trái) và perceptron AND sau huấn luyện (phải)")


def bang_dieu_kien(doc, w_an):
    cac_dong = []
    dk = {
        (2.0, 0.0, 1.0): "x₁ ≥ −0.5",
        (2.0, 0.0, -1.0): "x₁ ≥ 0.5",
        (1.0, 2.0, 1.0): "x₂ ≥ −0.5x₁ − 0.5",
        (1.0, 2.0, -1.0): "x₂ ≥ −0.5x₁ + 0.5",
        (1.0, 3.0, 0.0): "x₂ ≥ −x₁/3",
    }
    for k, w in enumerate(w_an):
        cac_dong.append([f"Nơ-ron {k + 1}", f"[{', '.join(so(v) for v in w)}]", f"v{chr(0x2081 + k)} = {btt(w, ['x₁', 'x₂', ''])}", f"o{chr(0x2081 + k)} = 1 ⇔ {dk[tuple(w)]}"])
    bang(doc, ["Nơ-ron ẩn", "Trọng số", "Tổng có trọng số", "Điều kiện kích hoạt"], cac_dong)


def viet_p2_cau1(doc):
    doc.add_heading("PHẦN 2 – CÂU 1. MIỀN KÍCH HOẠT CỦA MẠNG NƠ-RON", level=1).paragraph_format.page_break_before = True
    doan(doc, "Mạng gồm tầng vào (x₁, x₂ và đầu vào bias bằng 1), tầng ẩn 3 nơ-ron (1, 2, 3) và tầng ra 1 nơ-ron (4) có thêm đầu vào bias bằng 1. Mọi nơ-ron dùng hàm ngưỡng φ(v) = 1 nếu v ≥ 0, 0 nếu v < 0.")
    doan_nhieu_phan(doc, [("Quy ước thứ tự trọng số: ", True), ("đề không ghi thứ tự phần tử trong vector. Bài làm hiểu theo thứ tự trên hình: wⱼ = [trọng số của x₁, của x₂, của bias] cho nơ-ron ẩn j, và w₄ = [trọng số của o₁, o₂, o₃, của bias] cho nơ-ron ra.", False)])
    hinh(doc, p2c12.ve_mang(p2c12.W_AN, p2c12.W_RA, "p2_cau1_mang.png"), "Mạng nơ-ron 3 tầng với các vector trọng số đã cho", 14)

    doc.add_heading("1) Tổng có trọng số tại từng nơ-ron", level=2)
    for k, bt in enumerate(p2c12.bieu_thuc_an()):
        cong_thuc(doc, f"v{chr(0x2081 + k)} = {bt},      o{chr(0x2081 + k)} = φ(v{chr(0x2081 + k)})")
    cong_thuc(doc, f"v₄ = {p2c12.bieu_thuc_ra()},      y = φ(v₄)")
    doan(doc, "Nơ-ron 1 có trọng số của x₂ bằng 0 nên chỉ phụ thuộc x₁; nơ-ron 3 có bias bằng 0 nên ranh giới đi qua gốc tọa độ.")

    doc.add_heading("2) Điều kiện để mạng được kích hoạt (y = 1)", level=2)
    doan(doc, "Bước 1 – điều kiện kích hoạt của từng nơ-ron ẩn:")
    bang_dieu_kien(doc, p2c12.W_AN)
    doan(doc, "Bước 2 – xét nơ-ron ra. Vì oⱼ chỉ nhận 0 hoặc 1 nên chỉ có 8 tổ hợp (o₁, o₂, o₃):")
    bang(doc, ["(o₁, o₂, o₃)", "v₄ = 2o₁ + 3o₂ + o₃ + 1", "y = φ(v₄)"], [[f"({d['o'][0]}, {d['o'][1]}, {d['o'][2]})", so(d["v"]), d["y"]] for d in p2c12.bang_to_hop()])
    doan(doc, "Mọi trọng số của nơ-ron ra đều không âm và bias bằng +1 nên v₄ ≥ 1 > 0 với mọi tổ hợp. Kết luận: mạng luôn được kích hoạt, điều kiện y = 1 đúng với mọi (x₁, x₂) ∈ ℝ². Kết quả này được kiểm tra lại bằng chương trình trên lưới 601 × 601 điểm trong miền [−5, 4] × [−3, 4]: tỉ lệ điểm có y = 1 là " + f"{p2c12.khao_sat_luoi()['ty_le_kich_hoat'] * 100:g}%.")

    doc.add_heading("3) Biểu diễn miền kích hoạt trên mặt phẳng (x₁, x₂)", level=2)
    gd = p2c12.giao_diem()
    doan(doc, "Ba ranh giới của tầng ẩn: x₁ = −0.5 (nơ-ron 1), x₁ + 2x₂ + 1 = 0 (nơ-ron 2), x₁ + 3x₂ = 0 (nơ-ron 3). Giao điểm từng cặp: " + "; ".join(f"đường {i} và {j} tại ({so(round(p[0], 4))}, {so(round(p[1], 4))})" for i, j, p in gd) + ". Mũi tên chỉ phía nơ-ron tương ứng bằng 1; vùng tô xanh là miền y = 1.")
    hinh(doc, p2c12.ve_mot_mien(p2c12.W_AN, p2c12.W_RA, "Miền kích hoạt của mạng (Câu 1)", "p2_cau1_mien_kich_hoat.png"), "Ba ranh giới của tầng ẩn và miền kích hoạt của mạng (toàn bộ mặt phẳng)", 14)

    doc.add_heading("4) Ý nghĩa hình học của miền kích hoạt", level=2)
    mau = p2c12.khao_sat_luoi()["cac_mau_o"]
    gach_dau_dong(doc, "Mỗi nơ-ron ẩn là một perceptron: ranh giới vⱼ = 0 là một đường thẳng, chia mặt phẳng thành hai nửa mặt phẳng (oⱼ = 1 và oⱼ = 0).")
    gach_dau_dong(doc, f"Ba đường thẳng chia mặt phẳng thành {len(mau)} vùng, mỗi vùng ứng với một mẫu (o₁, o₂, o₃): " + ", ".join(f"({a}, {b}, {c})" for a, b, c in mau) + ". Tầng ẩn biến điểm (x₁, x₂) thành “mã vùng” nhị phân.")
    gach_dau_dong(doc, "Nơ-ron ra là perceptron trên mã vùng: nó chọn hợp một số vùng làm miền kích hoạt. Ở đây v₄ = 2o₁ + 3o₂ + o₃ + 1 dương với mọi mã vùng nên cả 7 vùng đều được chọn, miền kích hoạt là toàn bộ mặt phẳng.")
    gach_dau_dong(doc, "Nếu bias của nơ-ron ra âm đủ lớn thì miền kích hoạt mới bị thu hẹp. Ví dụ bias −5 thì chỉ còn (1, 1, 0) và (1, 1, 1) thỏa v₄ ≥ 0, tức giao của hai nửa mặt phẳng o₁ = 1 và o₂ = 1. Như vậy, mạng 2 tầng tạo được miền là hợp / giao của các nửa mặt phẳng (miền đa giác lồi hoặc ghép nhiều miền), điều một perceptron đơn không làm được.")


def viet_p2_cau2(doc):
    doc.add_heading("PHẦN 2 – CÂU 2. ẢNH HƯỞNG CỦA TRỌNG SỐ BIAS", level=1).paragraph_format.page_break_before = True
    doan(doc, "Thay các trọng số bias của tầng ẩn đang bằng 1 thành −1, các trọng số khác giữ nguyên (cùng quy ước thứ tự [x₁, x₂, bias]).")

    doc.add_heading("1) Vector trọng số sau khi đổi bias", level=2)
    cac_dong = []
    for k, (w_cu, w_moi) in enumerate(zip(p2c12.W_AN, p2c12.W_AN_MOI)):
        cac_dong.append([f"w{chr(0x2081 + k)}", f"[{', '.join(so(v) for v in w_cu)}]", f"[{', '.join(so(v) for v in w_moi)}]", "Đổi bias 1 → −1" if w_cu[2] != w_moi[2] else "Giữ nguyên (bias bằng 0, không phải 1)"])
    cac_dong.append(["w₄", f"[{', '.join(so(v) for v in p2c12.W_RA)}]", f"[{', '.join(so(v) for v in p2c12.W_RA)}]", "Giữ nguyên (tầng ra)"])
    bang(doc, ["Vector", "Trước", "Sau", "Ghi chú"], cac_dong)
    hinh(doc, p2c12.ve_mang(p2c12.W_AN_MOI, p2c12.W_RA, "p2_cau2_mang.png"), "Mạng nơ-ron sau khi đổi bias tầng ẩn", 14)

    doc.add_heading("2) Biểu thức tính toán và điều kiện kích hoạt mới", level=2)
    for k, bt in enumerate(p2c12.bieu_thuc_an(p2c12.W_AN_MOI)):
        cong_thuc(doc, f"v{chr(0x2081 + k)} = {bt},      o{chr(0x2081 + k)} = φ(v{chr(0x2081 + k)})")
    cong_thuc(doc, f"v₄ = {p2c12.bieu_thuc_ra()},      y = φ(v₄)")
    bang_dieu_kien(doc, p2c12.W_AN_MOI)

    doc.add_heading("3) Miền (x₁, x₂) làm đầu ra bằng 1", level=2)
    doan(doc, f"Nơ-ron ra không đổi nên bảng 8 tổ hợp ở Câu 1 vẫn đúng: v₄ ≥ 1 > 0 với mọi (o₁, o₂, o₃). Vì vậy đầu ra vẫn bằng 1 với mọi (x₁, x₂) ∈ ℝ² (kiểm tra trên lưới: {p2c12.khao_sat_luoi(p2c12.W_AN_MOI)['ty_le_kich_hoat'] * 100:g}% điểm có y = 1). Thay đổi chỉ xảy ra ở tầng ẩn: miền o₁ = 1 thành x₁ ≥ 0.5, miền o₂ = 1 thành x₂ ≥ −0.5x₁ + 0.5.")

    doc.add_heading("4) Miền kích hoạt mới", level=2)
    hinh(doc, p2c12.ve_mot_mien(p2c12.W_AN_MOI, p2c12.W_RA, "Miền kích hoạt sau khi đổi bias (Câu 2)", "p2_cau2_mien_kich_hoat.png"), "Ranh giới tầng ẩn và miền kích hoạt sau khi đổi bias", 14)

    doc.add_heading("5) So sánh trước và sau; vai trò của bias", level=2)
    hinh(doc, p2c12.ve_so_sanh(), "So sánh trước (trái) và sau (phải) khi đổi bias; đường chấm xám là ranh giới cũ")
    kc_cu, kc_moi = p2c12.khoang_cach_goc(p2c12.W_AN), p2c12.khoang_cach_goc(p2c12.W_AN_MOI)
    bang(
        doc,
        ["Nơ-ron", "Ranh giới trước", "Ranh giới sau", "Thay đổi"],
        [
            ["1", "x₁ = −0.5", "x₁ = 0.5", f"Tịnh tiến sang phải 1 đơn vị; khoảng cách tới gốc {so(round(kc_cu[0], 4))} → {so(round(kc_moi[0], 4))} nhưng gốc chuyển từ phía o₁ = 1 sang phía o₁ = 0"],
            ["2", "x₂ = −0.5x₁ − 0.5", "x₂ = −0.5x₁ + 0.5", f"Cùng hệ số góc −0.5, tịnh tiến song song lên 1 đơn vị; gốc chuyển sang phía o₂ = 0"],
            ["3", "x₂ = −x₁/3", "x₂ = −x₁/3", "Không đổi (bias bằng 0)"],
            ["4 (đầu ra)", "Toàn mặt phẳng", "Toàn mặt phẳng", "Không đổi"],
        ],
        chu_nho=True,
    )
    mau_cu, mau_moi = set(p2c12.khao_sat_luoi()["cac_mau_o"]), set(p2c12.khao_sat_luoi(p2c12.W_AN_MOI)["cac_mau_o"])
    gach_dau_dong(doc, "Bias không làm đổi hướng ranh giới (pháp tuyến vẫn là (w₁, w₂)) mà chỉ tịnh tiến song song ranh giới. Khoảng cách từ gốc tới ranh giới là |b| / ‖(w₁, w₂)‖ và dấu của b quyết định gốc tọa độ nằm ở phía kích hoạt (b > 0) hay không kích hoạt (b < 0).")
    gach_dau_dong(doc, "Đổi bias từ +1 sang −1 làm miền kích hoạt của nơ-ron 1 và 2 thu hẹp (nơ-ron “khó” kích hoạt hơn): cần x₁ ≥ 0.5 thay vì x₁ ≥ −0.5, và cần x₁ + 2x₂ ≥ 1 thay vì ≥ −1.")
    gach_dau_dong(doc, f"Cách chia vùng của tầng ẩn thay đổi: trước có mã vùng {', '.join(str(m) for m in sorted(mau_cu - mau_moi))} mà sau không còn, sau xuất hiện mã vùng mới {', '.join(str(m) for m in sorted(mau_moi - mau_cu))}; ba giao điểm đối xứng qua gốc tọa độ.")
    gach_dau_dong(doc, "Miền kích hoạt của cả mạng không đổi vì nơ-ron ra có bias dương và trọng số không âm, luôn kích hoạt. Bias ở tầng ẩn chỉ dời các ranh giới thành phần; muốn thay đổi miền kích hoạt cuối cùng phải chỉnh cả bias / trọng số của tầng ra.")


def viet_p2_cau3(doc):
    doc.add_heading("PHẦN 2 – CÂU 3. MẠNG NƠ-RON XÉT DANH HIỆU SINH VIÊN GIỎI TOÀN DIỆN", level=1).paragraph_format.page_break_before = True

    doc.add_heading("1) Biến đầu vào, đầu ra và vector trọng số", level=2)
    gach_dau_dong(doc, "Đầu vào: x₁ = điểm trung bình (thang 4, x₁ ∈ [0, 4]); x₂ = điểm hạnh kiểm (thang 10); x₃ = điểm rèn luyện (thang 10).")
    gach_dau_dong(doc, "Đầu ra: y ∈ {0, 1}, y = 1 là đạt danh hiệu, y = 0 là không đạt.")
    gach_dau_dong(doc, f"Vector trọng số theo hệ số của từng tiêu chí: w = (w₁, w₂, w₃) = ({', '.join(so(v) for v in p2c34.TRONG_SO)}).")

    doc.add_heading("2) Mạng perceptron phân loại", level=2)
    doan(doc, "Một perceptron 3 đầu vào (cộng đầu vào bias bằng 1) là đủ:")
    cong_thuc(doc, "v = w₁x₁ + w₂x₂ + w₃x₃ + b = 2x₁ + x₂ + x₃ + b = D + b,      y = f(v)")

    doc.add_heading("3) Bias và hàm kích hoạt", level=2)
    doan(doc, "Điều kiện đạt D ≥ 24 ⇔ 2x₁ + x₂ + x₃ − 24 ≥ 0. Chọn bias b = −24 và hàm bước:")
    cong_thuc(doc, "y = step(v) = 1 nếu v ≥ 0;  0 nếu v < 0,      v = 2x₁ + x₂ + x₃ − 24")
    doan(doc, f"Hàm bước dùng dấu “≥” để sinh viên có D đúng bằng 24 vẫn đạt, khớp với điều kiện D ≥ 24. Điểm D tối đa là 2·4 + 10 + 10 = {so(p2c34.diem_danh_hieu(p2c34.DIEM_TOI_DA))}, nên ngưỡng 24 là mức khá cao.")

    doc.add_heading("4) Sơ đồ kiến trúc mạng", level=2)
    hinh(doc, p2c34.ve_mang(), "Perceptron xét danh hiệu sinh viên giỏi toàn diện", 11)

    doc.add_heading("5) Đầu ra với ba sinh viên", level=2)
    kq = p2c34.phan_loai()
    bang(doc, ["Sinh viên", "Điểm TB x₁", "Hạnh kiểm x₂", "Rèn luyện x₃", "D = 2x₁ + x₂ + x₃", "v = D − 24", "y", "Kết quả"], [[d["ten"], so(d["x"][0]), so(d["x"][1]), so(d["x"][2]), f"2·{so(d['x'][0])} + {so(d['x'][1])} + {so(d['x'][2])} = {so(d['D'])}", so(d["v"]), d["y"], d["ket_qua"]] for d in kq], chu_nho=True)
    doan(doc, "SV1 và SV3 đạt danh hiệu; SV2 thiếu 1 điểm (D = 23) nên không đạt.")

    doc.add_heading("6) Vì sao một perceptron giải được bài toán", level=2)
    doan(doc, "Tập sinh viên đạt là {(x₁, x₂, x₃) : 2x₁ + x₂ + x₃ − 24 ≥ 0}, tức một nửa không gian trong ℝ³ giới hạn bởi mặt phẳng 2x₁ + x₂ + x₃ = 24. Hai lớp “đạt” và “không đạt” vì vậy khả tách tuyến tính, đúng loại ranh giới mà một perceptron biểu diễn được (siêu phẳng wᵀx + b = 0). Quy tắc xét là tổng có trọng số so với ngưỡng nên trọng số và bias lấy trực tiếp từ quy định, không cần huấn luyện. Bài toán cần nhiều tầng chỉ khi miền đạt không phải một nửa không gian, ví dụ khi thêm điều kiện “không có tiêu chí nào dưới mức tối thiểu”.")


def viet_p2_cau4(doc):
    doc.add_heading("PHẦN 2 – CÂU 4. SỰ THAY ĐỔI RANH GIỚI PHÂN LOẠI THEO BIAS", level=1).paragraph_format.page_break_before = True
    ss = p2c34.so_sanh_bias()

    doc.add_heading("1) Tổng có trọng số với bias mới", level=2)
    doan(doc, "Giữ nguyên trọng số đầu vào (2, 1, 1), thay bias −24 bằng một giá trị b tùy chọn:")
    cong_thuc(doc, "v = 2x₁ + x₂ + x₃ + b = D + b,      y = step(v)")

    doc.add_heading("2) Điều kiện phân loại theo bias mới", level=2)
    cong_thuc(doc, "y = 1   ⇔   D + b ≥ 0   ⇔   D ≥ −b")
    doan(doc, "Bias đóng vai trò ngưỡng điểm danh hiệu: ngưỡng mới bằng −b. Mặt phẳng ranh giới 2x₁ + x₂ + x₃ = −b song song với mặt phẳng cũ.")

    doc.add_heading("3) Ba giá trị bias và ngưỡng tương ứng", level=2)
    bang(doc, ["Bias b", "Ngưỡng D ≥ −b", "Ý nghĩa"], [[so(s["bias"]), f"D ≥ {so(s['nguong'])}", "Bias gốc của Câu 3" if s["bias"] == p2c34.BIAS else ("Nới lỏng tiêu chuẩn" if s["bias"] > p2c34.BIAS else ("Vượt điểm tối đa 28: không ai đạt" if s["nguong"] > 28 else "Siết chặt tiêu chuẩn"))] for s in ss])

    doc.add_heading("4) So sánh kết quả phân loại (dữ liệu Câu 3)", level=2)
    doan(doc, "Đề ghi “dữ liệu của Bài 1.6.3”, tức bảng ba sinh viên ở Câu 3 phía trên (D lần lượt 25, 23, 28).")
    cac_dong = []
    for s in ss:
        cac_dong.append([so(s["bias"]), so(s["nguong"])] + [f"{so(d['v'])} → {'Đạt' if d['y'] else 'Không'}" for d in s["kq"]] + [s["so_dat"]])
    bang(doc, ["Bias b", "Ngưỡng", "SV1 (D = 25): v, kết quả", "SV2 (D = 23): v, kết quả", "SV3 (D = 28): v, kết quả", "Số SV đạt"], cac_dong, chu_nho=True)
    hinh(doc, p2c34.ve_so_sanh_bias(), "Điểm danh hiệu của ba sinh viên và các ngưỡng tương ứng với từng bias")

    doc.add_heading("5) Tập sinh viên được chọn khi bias tăng hoặc giảm", level=2)
    gach_dau_dong(doc, "Bias tăng (−24 → −22): ngưỡng giảm còn 22, SV2 (D = 23) chuyển sang đạt; cả 3 sinh viên đạt.")
    gach_dau_dong(doc, "Bias giảm (−24 → −26): ngưỡng tăng lên 26, SV1 (D = 25) bị loại; chỉ còn SV3.")
    gach_dau_dong(doc, "Bias giảm tiếp (−29): ngưỡng 29 lớn hơn điểm tối đa 28 nên không sinh viên nào đạt, mạng luôn cho y = 0.")
    gach_dau_dong(doc, "Tập sinh viên được chọn lồng nhau: bias càng lớn, tập càng mở rộng (thêm người có D thấp hơn); bias càng nhỏ, tập càng thu hẹp. Thứ tự xếp hạng theo D không đổi, chỉ thay đổi vị trí điểm cắt.")

    doc.add_heading("6) Vai trò của bias", level=2)
    doan(doc, "Trọng số quyết định tầm quan trọng tương đối giữa các tiêu chí (hướng của mặt phẳng ranh giới), còn bias quyết định mức điểm tổng hợp tối thiểu để được chọn (vị trí của mặt phẳng). Đổi bias tương đương tịnh tiến song song ranh giới quyết định mà không đổi cách tính điểm, nên điều chỉnh bias chính là điều chỉnh độ khắt khe của tiêu chuẩn: b = −ngưỡng. Không có bias, ranh giới buộc phải đi qua gốc tọa độ (ngưỡng D ≥ 0) và mạng không biểu diễn được quy định D ≥ 24.")


def phu_luc(doc):
    doc.add_page_break()
    doc.add_heading("PHỤ LỤC. MÃ NGUỒN", level=1)
    doan(doc, f"Toàn bộ mã nguồn nằm trong thư mục {THU_MUC.name} trên nhánh {NHANH} của kho: {REPO}")
    bang(
        doc,
        ["File", "Nội dung"],
        [
            ["ve_hinh.py", "Hàm vẽ dùng chung: trục tọa độ, đường thẳng, tô miền, sơ đồ perceptron"],
            ["p1_cau1_ranh_gioi.py", "Phần 1 – Câu 1: perceptron cho ranh giới x₂ = −1.5x₁ + 3"],
            ["p1_cau2_perceptron_hinh_ve.py", "Phần 1 – Câu 2: ranh giới, đầu ra, perceptron theo hình, AND / OR"],
            ["p1_cau3_hoc_ham_or.py", "Phần 1 – Câu 3: perceptron learning (batch) cho hàm OR, ρ = 0.4"],
            ["p1_cau4_hoc_ham_and.py", "Phần 1 – Câu 4: huấn luyện perceptron AND theo từng epoch"],
            ["p2_cau1_2_mien_kich_hoat.py", "Phần 2 – Câu 1, 2: miền kích hoạt mạng 3 tầng, ảnh hưởng của bias"],
            ["p2_cau3_4_xet_danh_hieu.py", "Phần 2 – Câu 3, 4: perceptron xét danh hiệu, so sánh các bias"],
            ["tao_bao_cao.py", "Sinh file báo cáo Word này"],
        ],
    )
    doan(doc, "Cách chạy:")
    khoi_code(doc, "pip install -r requirements.txt\npython p1_cau1_ranh_gioi.py\npython p1_cau2_perceptron_hinh_ve.py\npython p1_cau3_hoc_ham_or.py\npython p1_cau4_hoc_ham_and.py\npython p2_cau1_2_mien_kich_hoat.py\npython p2_cau3_4_xet_danh_hieu.py\npython tao_bao_cao.py")


def main():
    doc = Document()
    cai_dat_tai_lieu(doc)
    them_so_trang(doc)
    trang_bia(doc)
    muc_luc(doc)
    viet_p1_cau1(doc)
    viet_p1_cau2(doc)
    viet_p1_cau3(doc)
    viet_p1_cau4(doc)
    viet_p2_cau1(doc)
    viet_p2_cau2(doc)
    viet_p2_cau3(doc)
    viet_p2_cau4(doc)
    phu_luc(doc)
    duong_dan = THU_MUC.parent.parent / TEN_FILE
    doc.save(duong_dan)
    print("Da tao:", duong_dan.name)


if __name__ == "__main__":
    main()
