import inspect
import shutil
from datetime import date
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import bai1_logistic_tre_so_sinh as bai1
import bai2_logistic_danh_gia_phim as bai2
import bai3_lan_truyen_nguoc as bai3

HO_TEN = "Võ Tường Long"
MA_SV = "3123580023"
TEN_FILE = "BT02_ComputationGraph_VoTuongLong_3123580023.docx"
REPO = "https://github.com/tuonglong1012/Mon_Hoc_Sau_DDU1231/tree/BT02_ComputationGraph"
THU_MUC = Path(__file__).resolve().parent
MAU_TIEU_DE = RGBColor(0x0D, 0x47, 0xA1)


def f(x, n=6):
    return f"{x:.{n}f}"


def to_mau_o(o, mau):
    tc_pr = o._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), mau)
    tc_pr.append(shd)


def cai_dat_tai_lieu(doc):
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
    kieu = doc.styles["Normal"]
    kieu.font.name = "Times New Roman"
    kieu.font.size = Pt(13)
    kieu.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    kieu.paragraph_format.space_after = Pt(4)
    kieu.paragraph_format.line_spacing = 1.2
    for ten, co in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 13)):
        h = doc.styles[ten]
        h.font.name = "Times New Roman"
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        h.font.size = Pt(co)
        h.font.bold = True
        h.font.color.rgb = MAU_TIEU_DE
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)


def them_so_trang(doc):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for loai, noi_dung in (("begin", None), (None, "PAGE"), ("end", None)):
        if loai:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), loai)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = noi_dung
        run._r.append(el)


def doan(doc, van_ban, dam=False, nghieng=False, can=None, co=None):
    p = doc.add_paragraph()
    run = p.add_run(van_ban)
    run.bold = dam
    run.italic = nghieng
    if co:
        run.font.size = Pt(co)
    if can:
        p.alignment = can
    return p


def doan_nhieu_phan(doc, cac_phan):
    p = doc.add_paragraph()
    for noi_dung, dam in cac_phan:
        r = p.add_run(noi_dung)
        r.bold = dam
    return p


def cong_thuc(doc, van_ban):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(van_ban)
    run.font.name = "Cambria Math"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria Math")
    run.font.size = Pt(13)
    return p


def gach_dau_dong(doc, van_ban):
    return doc.add_paragraph(van_ban, style="List Bullet")


def bang(doc, tieu_de_cot, cac_dong, do_rong=None, chu_nho=False):
    t = doc.add_table(rows=1, cols=len(tieu_de_cot))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, ten in enumerate(tieu_de_cot):
        o = t.rows[0].cells[i]
        o.text = ""
        r = o.paragraphs[0].add_run(ten)
        r.bold = True
        r.font.size = Pt(10 if chu_nho else 11)
        o.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        to_mau_o(o, "DCE6F1")
    for dong in cac_dong:
        cells = t.add_row().cells
        for i, gt in enumerate(dong):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(gt))
            r.font.size = Pt(10 if chu_nho else 11)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if do_rong:
        for dong in t.rows:
            for i, w in enumerate(do_rong):
                dong.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def hinh(doc, duong_dan, chu_thich, rong=16):
    doc.add_picture(str(duong_dan), width=Cm(rong))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doan(doc, chu_thich, nghieng=True, can=WD_ALIGN_PARAGRAPH.CENTER, co=11)


def khoi_code(doc, ma):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    o = t.rows[0].cells[0]
    to_mau_o(o, "F5F5F5")
    o.text = ""
    p = o.paragraphs[0]
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(0)
    dong = ma.rstrip("\n").split("\n")
    for k, d in enumerate(dong):
        r = p.add_run(d)
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        r.font.size = Pt(9)
        if k < len(dong) - 1:
            r.add_break(WD_BREAK.LINE)
    doc.add_paragraph()


def trang_bia(doc):
    for _ in range(3):
        doc.add_paragraph()
    doan(doc, "BÁO CÁO BÀI TẬP", dam=True, can=WD_ALIGN_PARAGRAPH.CENTER, co=22)
    doan(doc, "MÔN HỌC: HỌC SÂU (DEEP LEARNING)", dam=True, can=WD_ALIGN_PARAGRAPH.CENTER, co=16)
    doc.add_paragraph()
    doan(doc, "BÀI TẬP 02", dam=True, can=WD_ALIGN_PARAGRAPH.CENTER, co=18)
    doan(doc, "HỒI QUY LOGISTIC VÀ ĐỒ THỊ TÍNH TOÁN", dam=True, can=WD_ALIGN_PARAGRAPH.CENTER, co=20).runs[0].font.color.rgb = MAU_TIEU_DE
    doan(doc, "(Logistic Regression & Computation Graph)", nghieng=True, can=WD_ALIGN_PARAGRAPH.CENTER, co=14)
    for _ in range(5):
        doc.add_paragraph()
    t = doc.add_table(rows=3, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (nhan, gt) in enumerate((("Họ và tên sinh viên:", HO_TEN), ("Mã số sinh viên:", MA_SV), ("Mã nguồn:", "github.com/tuonglong1012/Mon_Hoc_Sau_DDU1231\n(nhánh BT02_ComputationGraph)"))):
        t.rows[i].cells[0].text = ""
        r = t.rows[i].cells[0].paragraphs[0].add_run(nhan)
        r.bold = True
        r.font.size = Pt(14)
        t.rows[i].cells[1].text = ""
        r = t.rows[i].cells[1].paragraphs[0].add_run(gt)
        r.font.size = Pt(14)
        t.rows[i].cells[0].width = Cm(5.5)
        t.rows[i].cells[1].width = Cm(10)
    for _ in range(8):
        doc.add_paragraph()
    doan(doc, f"Tháng {date.today().month} năm {date.today().year}", nghieng=True, can=WD_ALIGN_PARAGRAPH.CENTER, co=13)
    doc.add_page_break()


def chung_minh_gradient(doc, tham_so_cat, tham_so_trong, bien_x):
    doan(doc, "Xét một mẫu (xᵢ, yᵢ). Theo đồ thị tính toán, hàm mất mát của mẫu là hàm hợp:", nghieng=False)
    cong_thuc(doc, f"zᵢ = {tham_so_trong}·{bien_x} + {tham_so_cat}   →   ŷᵢ = σ(zᵢ)   →   Lᵢ = −[yᵢ ln ŷᵢ + (1 − yᵢ) ln(1 − ŷᵢ)]")
    doan_nhieu_phan(doc, [("Bước 1 – Đạo hàm của hàm sigmoid: ", True), ("σ(z) = (1 + e⁻ᶻ)⁻¹ nên", False)])
    cong_thuc(doc, "σ′(z) = e⁻ᶻ / (1 + e⁻ᶻ)² = [1/(1 + e⁻ᶻ)]·[e⁻ᶻ/(1 + e⁻ᶻ)] = σ(z)·(1 − σ(z))")
    doan(doc, "Do đó ∂ŷᵢ/∂zᵢ = ŷᵢ(1 − ŷᵢ).")
    doan_nhieu_phan(doc, [("Bước 2 – Đạo hàm của BCE theo ŷᵢ:", True)])
    cong_thuc(doc, "∂Lᵢ/∂ŷᵢ = −yᵢ/ŷᵢ + (1 − yᵢ)/(1 − ŷᵢ) = [−yᵢ(1 − ŷᵢ) + (1 − yᵢ)ŷᵢ] / [ŷᵢ(1 − ŷᵢ)] = (ŷᵢ − yᵢ) / [ŷᵢ(1 − ŷᵢ)]")
    doan_nhieu_phan(doc, [("Bước 3 – Quy tắc dây chuyền theo zᵢ:", True)])
    cong_thuc(doc, "∂Lᵢ/∂zᵢ = (∂Lᵢ/∂ŷᵢ)·(∂ŷᵢ/∂zᵢ) = (ŷᵢ − yᵢ) / [ŷᵢ(1 − ŷᵢ)] · ŷᵢ(1 − ŷᵢ) = ŷᵢ − yᵢ")
    doan_nhieu_phan(doc, [("Bước 4 – Đạo hàm theo tham số: ", True), (f"vì ∂zᵢ/∂{tham_so_cat} = 1 và ∂zᵢ/∂{tham_so_trong} = {bien_x} nên", False)])
    cong_thuc(doc, f"∂Lᵢ/∂{tham_so_cat} = (ŷᵢ − yᵢ)·1 = ŷᵢ − yᵢ,      ∂Lᵢ/∂{tham_so_trong} = (ŷᵢ − yᵢ)·{bien_x}")
    doan_nhieu_phan(doc, [("Bước 5 – Lấy trung bình trên m mẫu: ", True), ("vì J = (1/m)·Σ Lᵢ và đạo hàm là toán tử tuyến tính nên", False)])
    cong_thuc(doc, f"∂J/∂{tham_so_cat} = (1/m)·Σᵢ (ŷᵢ − yᵢ),      ∂J/∂{tham_so_trong} = (1/m)·Σᵢ (ŷᵢ − yᵢ)·{bien_x}      (đpcm)")


def bang_buoc_bai1(doc, buoc):
    cac_dong = []
    for i in range(len(bai1.X)):
        cac_dong.append([i + 1, bai1.X[i], int(bai1.Y[i]), f(buoc["z"][i]), f(buoc["y_hat"][i]), f(buoc["sai_so"][i]), f(buoc["sai_so_x"][i])])
    cac_dong.append(["Σ", "", "", "", "", f(buoc["sai_so"].sum()), f(buoc["sai_so_x"].sum())])
    bang(doc, ["i", "xᵢ", "yᵢ", "zᵢ", "ŷᵢ", "ŷᵢ − yᵢ", "(ŷᵢ − yᵢ)xᵢ"], cac_dong)


def viet_bai1(doc):
    ls = bai1.huan_luyen(2)
    b0, b1, b2 = ls
    doc.add_heading("BÀI 1. HỒI QUY LOGISTIC PHÂN LOẠI GIỚI TÍNH TRẺ SƠ SINH", level=1)
    doan(doc, "Dữ liệu: x = (2.4; 3.3; 3.2; 2.5) kg, y = (0; 1; 1; 0) với 0 là Nữ, 1 là Nam (đề ghi nhầm “0 là nam”, theo bảng dữ liệu thì Nam có nhãn 1). Mô hình ŷ = σ(θ₀ + θ₁x), m = 4, khởi tạo θ₀⁽⁰⁾ = θ₁⁽⁰⁾ = 0.5, α = 0.2.")

    doc.add_heading("a) Đầu vào tuyến tính zᵢ và xác suất dự đoán ŷᵢ", level=2)
    cong_thuc(doc, "zᵢ = θ₀ + θ₁xᵢ,        ŷᵢ = h_θ(xᵢ) = σ(zᵢ) = 1 / (1 + e^(−zᵢ))")
    doan(doc, "Với θ₀ = θ₁ = 0.5: zᵢ = 0.5 + 0.5xᵢ. Ví dụ mẫu 1: z₁ = 0.5 + 0.5·2.4 = 1.7, ŷ₁ = 1/(1 + e^(−1.7)) = 0.845535.")
    bang(doc, ["i", "xᵢ (kg)", "yᵢ", "zᵢ = 0.5 + 0.5xᵢ", "ŷᵢ = σ(zᵢ)"], [[i + 1, bai1.X[i], int(bai1.Y[i]), f(b0["z"][i]), f(b0["y_hat"][i])] for i in range(4)])
    doan(doc, "Nhận xét: cả 4 mẫu đều có ŷᵢ > 0.5, tức mô hình ban đầu dự đoán tất cả là Nam, sai ở hai bé gái (mẫu 1 và 4).")

    doc.add_heading("b) Binary Cross-Entropy Loss tại tham số khởi tạo", level=2)
    cong_thuc(doc, "Lᵢ = −[yᵢ ln(ŷᵢ) + (1 − yᵢ) ln(1 − ŷᵢ)]")
    doan(doc, "Với yᵢ = 0 thì Lᵢ = −ln(1 − ŷᵢ); với yᵢ = 1 thì Lᵢ = −ln(ŷᵢ).")
    cac_dong = []
    for i in range(4):
        bt = f"−ln(1 − {f(b0['y_hat'][i])})" if bai1.Y[i] == 0 else f"−ln({f(b0['y_hat'][i])})"
        cac_dong.append([i + 1, int(bai1.Y[i]), f(b0["y_hat"][i]), bt, f(b0["loss"][i])])
    cac_dong.append(["", "", "", "Tổng", f(b0["loss"].sum())])
    bang(doc, ["i", "yᵢ", "ŷᵢ", "Biểu thức", "Lᵢ"], cac_dong)
    cong_thuc(doc, f"J(θ₀⁽⁰⁾, θ₁⁽⁰⁾) = (1/4)·{f(b0['loss'].sum())} = {f(b0['J'])}")

    doc.add_heading("c) Chứng minh công thức đạo hàm", level=2)
    chung_minh_gradient(doc, "θ₀", "θ₁", "xᵢ")

    doc.add_heading("d) Gradient tại thời điểm khởi tạo và lần cập nhật thứ nhất", level=2)
    bang_buoc_bai1(doc, b0)
    cong_thuc(doc, f"∂J/∂θ₀ = {f(b0['sai_so'].sum())} / 4 = {f(b0['grad0'])},      ∂J/∂θ₁ = {f(b0['sai_so_x'].sum())} / 4 = {f(b0['grad1'])}")
    cong_thuc(doc, f"θ₀⁽¹⁾ = 0.5 − 0.2·{f(b0['grad0'])} = {f(b1['theta0'])}")
    cong_thuc(doc, f"θ₁⁽¹⁾ = 0.5 − 0.2·{f(b0['grad1'])} = {f(b1['theta1'])}")

    doc.add_heading("e) Lần cập nhật thứ hai", level=2)
    doan(doc, f"Với θ₀⁽¹⁾ = {f(b1['theta0'])}, θ₁⁽¹⁾ = {f(b1['theta1'])}: zᵢ = {f(b1['theta0'])} + {f(b1['theta1'])}·xᵢ.")
    bang_buoc_bai1(doc, b1)
    cong_thuc(doc, f"∂J/∂θ₀ = {f(b1['sai_so'].sum())} / 4 = {f(b1['grad0'])},      ∂J/∂θ₁ = {f(b1['sai_so_x'].sum())} / 4 = {f(b1['grad1'])}")
    cong_thuc(doc, f"θ₀⁽²⁾ = {f(b1['theta0'])} − 0.2·{f(b1['grad0'])} = {f(b2['theta0'])}")
    cong_thuc(doc, f"θ₁⁽²⁾ = {f(b1['theta1'])} − 0.2·{f(b1['grad1'])} = {f(b2['theta1'])}")

    doc.add_heading("f) Hàm mất mát tại ba thời điểm", level=2)
    bang(
        doc,
        ["Thời điểm", "θ₀", "θ₁", "L₁", "L₂", "L₃", "L₄", "J"],
        [[f"t = {b['t']}", f(b["theta0"], 4), f(b["theta1"], 4)] + [f(v, 4) for v in b["loss"]] + [f(b["J"])] for b in ls],
        chu_nho=True,
    )
    doan(doc, f"Nhận xét: J giảm đơn điệu {f(b0['J'], 4)} → {f(b1['J'], 4)} → {f(b2['J'], 4)} (giảm {100 * (b0['J'] - b1['J']) / b0['J']:.1f}% ở lần 1 và {100 * (b1['J'] - b2['J']) / b1['J']:.1f}% ở lần 2), chứng tỏ Gradient Descent đi đúng hướng ngược gradient. Loss của hai bé gái (mẫu 1, 4) giảm mạnh, còn loss của hai bé trai (mẫu 2, 3) tăng nhẹ vì θ₁ giảm làm mọi xác suất cùng giảm. Mức giảm ở lần 2 nhỏ hơn lần 1 vì gradient đã nhỏ dần. Sau 2 bước J vẫn còn lớn (≈ 0.76) do đặc trưng x không được chuẩn hóa và cần nhiều vòng lặp hơn để hội tụ.")

    doc.add_heading("g) Dự đoán với mô hình sau hai lần cập nhật", level=2)
    doan(doc, f"Mô hình: ŷ = σ({f(b2['theta0'])} + {f(b2['theta1'])}·x), ngưỡng 0.5.")
    kq = bai1.du_doan(b2["theta0"], b2["theta1"])
    bang(doc, ["x (kg)", "z", "P(Nam) = ŷ", "So với 0.5", "Kết luận"], [[k["x"], f(k["z"]), f(k["p"]), "≥ 0.5" if k["nhan"] else "< 0.5", k["gioi_tinh"]] for k in kq])
    doan(doc, "Cả hai trẻ đều được dự đoán là Nam. Kết quả với x = 3.5 kg là hợp lý (nặng hơn các bé trai trong dữ liệu). Với x = 2.8 kg (nằm giữa hai nhóm), mô hình sau chỉ 2 bước vẫn dự đoán mọi mẫu là Nam vì z = θ₀ + θ₁x > 0 với mọi x > 0, nên chưa phân biệt được tốt; cần huấn luyện thêm để mô hình học được ranh giới nằm giữa 2.5 kg và 3.2 kg.")

    doc.add_heading("h) Đồ thị tính toán cho một mẫu huấn luyện", level=2)
    hinh(doc, bai1.ve_do_thi(ls), "Hình 1. Đồ thị tính toán của Logistic Regression (mẫu 1, tham số khởi tạo)")
    doan_nhieu_phan(doc, [("Lan truyền tiến (mũi tên liền, trái → phải): ", True), ("u = θ₁·x = 1.2 → z = u + θ₀ = 1.7 → ŷ = σ(z) = 0.8455 → L = −[y ln ŷ + (1 − y) ln(1 − ŷ)] = 1.8678.", False)])
    doan_nhieu_phan(doc, [("Lan truyền ngược (mũi tên đứt, phải → trái): ", True), ("bắt đầu từ ∂L/∂L = 1; ∂L/∂ŷ = −y/ŷ + (1 − y)/(1 − ŷ) = 6.4739; ∂L/∂z = ∂L/∂ŷ·ŷ(1 − ŷ) = ŷ − y = 0.8455; ∂L/∂θ₀ = ∂L/∂z·1 = 0.8455; ∂L/∂u = 0.8455; ∂L/∂θ₁ = ∂L/∂u·x = 0.8455·2.4 = 2.0293. Lấy trung bình các gradient này trên 4 mẫu ta được đúng ∂J/∂θ₀, ∂J/∂θ₁ ở câu (d).", False)])
    doc.add_page_break()


def bang_buoc_bai2(doc, buoc):
    cac_dong = []
    for i in range(len(bai2.X)):
        cac_dong.append([i + 1, int(bai2.X[i]), f(bai2.X_CHUAN[i], 2), int(bai2.Y[i]), f(buoc["z"][i]), f(buoc["y_hat"][i]), f(buoc["loss"][i]), f(buoc["sai_so"][i]), f(buoc["sai_so_x"][i])])
    cac_dong.append(["Σ", "", "", "", "", "", f(buoc["loss"].sum()), f(buoc["sai_so"].sum()), f(buoc["sai_so_x"].sum())])
    bang(doc, ["i", "xᵢ", "xᵢ′", "yᵢ", "zᵢ", "ŷᵢ", "Lᵢ", "ŷᵢ − yᵢ", "(ŷᵢ − yᵢ)xᵢ′"], cac_dong, chu_nho=True)


def viet_bai2(doc):
    ls = bai2.huan_luyen(2)
    b0, b1, b2 = ls
    w2, bb2 = b2["w"], b2["b"]
    doc.add_heading("BÀI 2. HỒI QUY LOGISTIC DỰ ĐOÁN ĐÁNH GIÁ PHIM", level=1)

    doc.add_heading("a) Đầu vào, đầu ra, tham số và biểu diễn dữ liệu", level=2)
    gach_dau_dong(doc, "Đầu vào: thời lượng phim xᵢ (phút), khi huấn luyện dùng đặc trưng chuẩn hóa xᵢ′ = (xᵢ − 130)/50.")
    gach_dau_dong(doc, "Đầu ra: ŷᵢ = σ(w·xᵢ′ + b) ∈ (0, 1) là xác suất phim được đánh giá Tốt; nhãn dự đoán là 1 (G) nếu ŷᵢ ≥ 0.5, ngược lại là 0 (B).")
    gach_dau_dong(doc, "Tham số cần học: trọng số w và độ lệch b (khởi tạo w⁽⁰⁾ = b⁽⁰⁾ = 0, α = 0.2, m = 14).")
    hinh(doc, bai2.ve_du_lieu(), "Hình 2. 14 bộ phim trên hệ trục (thời lượng, nhãn đánh giá)", rong=15)
    doan(doc, "Quan sát: mọi phim Không tốt có thời lượng ≤ 119 phút, mọi phim Tốt có thời lượng ≥ 132 phút, nên dữ liệu tách được tuyến tính với ngưỡng nằm trong khoảng (119; 132).")

    doc.add_heading("b) Đặc trưng chuẩn hóa, zᵢ và ŷᵢ tại thời điểm khởi tạo", level=2)
    doan(doc, "Với w = b = 0 thì zᵢ = 0·xᵢ′ + 0 = 0 và ŷᵢ = σ(0) = 0.5 với mọi mẫu.")
    bang(doc, ["i", "xᵢ", "xᵢ′ = (xᵢ − 130)/50", "yᵢ", "zᵢ", "ŷᵢ"], [[i + 1, int(bai2.X[i]), f(bai2.X_CHUAN[i], 2), int(bai2.Y[i]), f(b0["z"][i], 1), f(b0["y_hat"][i], 1)] for i in range(14)])

    doc.add_heading("c) BCE Loss của từng mẫu và loss trung bình", level=2)
    doan(doc, "Vì ŷᵢ = 0.5 với mọi i nên: nếu yᵢ = 1 thì Lᵢ = −ln 0.5 = ln 2; nếu yᵢ = 0 thì Lᵢ = −ln(1 − 0.5) = ln 2. Vậy Lᵢ = ln 2 ≈ 0.693147 với cả 14 mẫu.")
    bang(doc, ["i", "yᵢ", "ŷᵢ", "Lᵢ"], [[i + 1, int(bai2.Y[i]), "0.5", f(b0["loss"][i])] for i in range(14)])
    cong_thuc(doc, f"J(0, 0) = (1/14)·14·ln 2 = ln 2 ≈ {f(b0['J'])}")

    doc.add_heading("d) Chứng minh biểu thức gradient", level=2)
    chung_minh_gradient(doc, "b", "w", "xᵢ′")

    doc.add_heading("e) Hai lần cập nhật Batch Gradient Descent", level=2)
    doc.add_heading("Lần cập nhật thứ nhất (t = 0 → 1)", level=3)
    bang_buoc_bai2(doc, b0)
    cong_thuc(doc, f"∂J/∂w = {f(b0['sai_so_x'].sum())}/14 = {f(b0['grad_w'])},      ∂J/∂b = {f(b0['sai_so'].sum())}/14 = {f(b0['grad_b'])}")
    doan(doc, "∂J/∂b = 0 vì có 7 phim tốt (ŷ − y = −0.5) và 7 phim không tốt (ŷ − y = +0.5) triệt tiêu nhau.")
    cong_thuc(doc, f"w⁽¹⁾ = 0 − 0.2·({f(b0['grad_w'])}) = {f(b1['w'])},      b⁽¹⁾ = 0 − 0.2·{f(b0['grad_b'])} = {f(b1['b'])}")
    doc.add_heading("Lần cập nhật thứ hai (t = 1 → 2)", level=3)
    bang_buoc_bai2(doc, b1)
    doan(doc, f"Loss sau lần cập nhật thứ nhất: J(w⁽¹⁾, b⁽¹⁾) = {f(b1['loss'].sum())}/14 = {f(b1['J'])}.")
    cong_thuc(doc, f"∂J/∂w = {f(b1['sai_so_x'].sum())}/14 = {f(b1['grad_w'])},      ∂J/∂b = {f(b1['sai_so'].sum())}/14 = {f(b1['grad_b'])}")
    cong_thuc(doc, f"w⁽²⁾ = {f(b1['w'])} − 0.2·({f(b1['grad_w'])}) = {f(w2)}")
    cong_thuc(doc, f"b⁽²⁾ = {f(b1['b'])} − 0.2·({f(b1['grad_b'])}) = {f(bb2)}")
    doc.add_heading("Loss sau lần cập nhật thứ hai", level=3)
    bang(doc, ["i", "zᵢ", "ŷᵢ", "Lᵢ"], [[i + 1, f(b2["z"][i]), f(b2["y_hat"][i]), f(b2["loss"][i])] for i in range(14)])
    bang(doc, ["Thời điểm", "w", "b", "J(w, b)"], [[f"t = {b['t']}", f(b["w"]), f(b["b"]), f(b["J"])] for b in ls])
    doan(doc, "Hàm mất mát giảm dần 0.693147 → 0.677941 → 0.663341, mỗi bước giảm khoảng 0.015.")

    doc.add_heading("f) Phương trình mô hình sau lần cập nhật thứ hai và dự đoán", level=2)
    he_so_x = w2 / bai2.DO_RONG
    he_so_tu_do = bb2 - w2 * bai2.TRUNG_TAM / bai2.DO_RONG
    cong_thuc(doc, f"ŷ = σ({f(w2)}·(x − 130)/50 + {f(bb2)}) = σ({f(he_so_x, 6)}·x − {f(-he_so_tu_do, 6)})")
    kq = bai2.du_doan(w2, bb2)
    bang(doc, ["x (phút)", "x′", "z = w·x′ + b", "P(Tốt) = ŷ", "Nhãn (ngưỡng 0.5)"], [[f"{k['x']:.0f}", f(k["x_chuan"], 2), f(k["z"]), f(k["p"]), "G (Tốt)" if k["nhan"] else "B (Không tốt)"] for k in kq])
    doan(doc, "Phim 125 phút có xác suất được đánh giá tốt ≈ 49.73% (dự đoán Không tốt), phim 160 phút ≈ 51.65% (dự đoán Tốt). Các xác suất còn rất gần 0.5 vì sau 2 bước w vẫn nhỏ, mô hình chưa “tự tin”.")

    doc.add_heading("g) Nhãn dự đoán, ma trận nhầm lẫn và các độ đo", level=2)
    dg = bai2.danh_gia(w2, bb2)
    bang(doc, ["i", "xᵢ", "ŷᵢ", "Nhãn dự đoán", "Nhãn thực tế", "Đúng/Sai"], [[i + 1, int(bai2.X[i]), f(dg["p"][i]), int(dg["y_du_doan"][i]), int(bai2.Y[i]), "Đúng" if dg["y_du_doan"][i] == bai2.Y[i] else "Sai"] for i in range(14)])
    bang(doc, ["", "Dự đoán G (1)", "Dự đoán B (0)"], [["Thực tế G (1)", f"TP = {dg['TP']}", f"FN = {dg['FN']}"], ["Thực tế B (0)", f"FP = {dg['FP']}", f"TN = {dg['TN']}"]])
    cong_thuc(doc, f"Accuracy = (TP + TN)/14 = ({dg['TP']} + {dg['TN']})/14 = {dg['accuracy']:.4f}")
    cong_thuc(doc, f"Precision = TP/(TP + FP) = {dg['TP']}/{dg['TP'] + dg['FP']} = {dg['precision']:.4f}")
    cong_thuc(doc, f"Recall = TP/(TP + FN) = {dg['TP']}/{dg['TP'] + dg['FN']} = {dg['recall']:.4f}")
    cong_thuc(doc, f"F1 = 2·P·R/(P + R) = {dg['f1']:.4f}")
    doan(doc, "Dù xác suất dự đoán còn gần 0.5, dấu của zᵢ đã khớp hoàn toàn với nhãn nên mô hình phân loại đúng cả 14 phim. Lý do là sau chuẩn hóa, các phim tốt có x′ > 0 và các phim không tốt có x′ < 0, còn b ≈ 0 nên ranh giới gần x = 130 phút.")

    doc.add_heading("h) Đường ranh giới quyết định", level=2)
    xc, xr = bai2.ranh_gioi_quyet_dinh(w2, bb2)
    doan(doc, "ŷ = 0.5 ⇔ σ(z) = 0.5 ⇔ z = 0 ⇔ w·x′ + b = 0. Do đó:")
    cong_thuc(doc, f"x′ = −b/w = −{f(bb2)}/{f(w2)} = {f(xc)}   ⇔   x = 130 + 50·x′ = {xr:.4f} phút")
    doan(doc, f"Phim dài hơn khoảng {xr:.2f} phút được dự đoán là Tốt, ngắn hơn là Không tốt.")
    hinh(doc, bai2.ve_sigmoid([("Sau 2 lần cập nhật", w2, bb2)], "bai2h_sigmoid.png", "Bài 2h - Đường Sigmoid và ranh giới quyết định sau 2 lần cập nhật"), "Hình 3. Đường sigmoid sau 2 lần cập nhật, ngưỡng 0.5 và ranh giới quyết định")
    doan(doc, "Đường sigmoid còn rất thoải (w nhỏ) nên xác suất chỉ dao động trong khoảng 0.46 – 0.54 trên miền dữ liệu.")

    doc.add_heading("i) Đồ thị tính toán, Forward Pass và Backward Pass", level=2)
    hinh(doc, bai2.ve_do_thi(ls), "Hình 4. Đồ thị tính toán cho mẫu 1 (x = 132, y = 1) tại w = b = 0")
    doan_nhieu_phan(doc, [("Forward Pass: ", True), ("x′ = (132 − 130)/50 = 0.04; u = w·x′ = 0; z = u + b = 0; ŷ = σ(0) = 0.5; L = −ln 0.5 = 0.6931.", False)])
    doan_nhieu_phan(doc, [("Backward Pass: ", True), ("∂L/∂L = 1 → ∂L/∂ŷ = −y/ŷ + (1 − y)/(1 − ŷ) = −1/0.5 = −2 → ∂L/∂z = ∂L/∂ŷ·ŷ(1 − ŷ) = −2·0.25 = −0.5 = ŷ − y → ∂L/∂b = ∂L/∂z·1 = −0.5, ∂L/∂u = −0.5 → ∂L/∂w = ∂L/∂u·x′ = −0.5·0.04 = −0.02.", False)])
    doan(doc, "Thực hiện tương tự cho 14 mẫu rồi lấy trung bình ta được ∂J/∂w = −0.277143 và ∂J/∂b = 0 như ở câu (e).")

    doc.add_heading("j) Huấn luyện 1000 lần lặp bằng Python (NumPy và PyTorch)", level=2)
    w_np, b_np, loss_np = bai2.huan_luyen_nhanh()
    w_pt, b_pt, loss_pt = bai2.huan_luyen_pytorch()
    doan(doc, "Mã nguồn đầy đủ ở file bai2_logistic_danh_gia_phim.py. Hai hàm huấn luyện chính:")
    khoi_code(doc, inspect.getsource(bai2.gradient) + "\n\n" + inspect.getsource(bai2.huan_luyen_nhanh) + "\n\n" + inspect.getsource(bai2.huan_luyen_pytorch))
    hinh(doc, bai2.ve_loss(loss_np, loss_pt, ls), "Hình 5. Hàm mất mát theo số lần lặp (NumPy và PyTorch trùng nhau)")
    dg_np = bai2.danh_gia(w_np, b_np)
    kq_np = bai2.du_doan(w_np, b_np)
    _, xr_np = bai2.ranh_gioi_quyet_dinh(w_np, b_np)
    bang(
        doc,
        ["Mô hình", "w", "b", "J", "Ranh giới x (phút)", "P(Tốt | 125)", "P(Tốt | 160)", "Accuracy"],
        [
            ["Sau 2 lần (câu e)", f(w2, 4), f(bb2, 4), f(b2["J"], 4), f"{xr:.2f}", f(kq[0]["p"], 4), f(kq[1]["p"], 4), f"{dg['accuracy']:.2f}"],
            ["1000 lần (NumPy)", f(w_np, 4), f(b_np, 4), f(loss_np[-1], 4), f"{xr_np:.2f}", f(kq_np[0]["p"], 4), f(kq_np[1]["p"], 4), f"{dg_np['accuracy']:.2f}"],
            ["1000 lần (PyTorch)", f(w_pt, 4), f(b_pt, 4), f(loss_pt[-1], 4), f"{bai2.ranh_gioi_quyet_dinh(w_pt, b_pt)[1]:.2f}", "", "", ""],
        ],
        chu_nho=True,
    )
    hinh(doc, bai2.ve_sigmoid([("Sau 2 lần cập nhật", w2, bb2), ("Sau 1000 lần lặp", w_np, b_np)], "bai2j_so_sanh_sigmoid.png", "Bài 2j - So sánh mô hình sau 2 lần cập nhật và sau 1000 lần lặp"), "Hình 6. So sánh đường sigmoid sau 2 lần cập nhật và sau 1000 lần lặp")
    doan(doc, "Nhận xét:")
    gach_dau_dong(doc, f"Ba giá trị loss đầu tiên của chương trình trùng khớp với tính tay ở câu (e) (0.693147; 0.677941; 0.663341). Kết quả NumPy và PyTorch giống nhau đến 6 chữ số thập phân.")
    gach_dau_dong(doc, f"Sau 1000 lần lặp loss giảm từ 0.6931 xuống {f(loss_np[-1], 4)}; w tăng từ {f(w2, 4)} lên {f(w_np, 4)} nên đường sigmoid dốc hơn nhiều, mô hình tự tin hơn: phim 160 phút có P(Tốt) ≈ {f(kq_np[1]['p'], 4)}.")
    gach_dau_dong(doc, f"Ranh giới quyết định dịch từ {xr:.2f} phút về {xr_np:.2f} phút, nằm giữa phim không tốt dài nhất (119 phút) và phim tốt ngắn nhất (132 phút). Phim 125 phút giờ có P(Tốt) ≈ {f(kq_np[0]['p'], 4)}, nằm ngay trên ranh giới.")
    gach_dau_dong(doc, "Accuracy vẫn bằng 1 vì dữ liệu tách được tuyến tính; loss vẫn tiếp tục giảm (về 0 khi w → ∞) nếu lặp thêm, do không dùng regularization.")
    doc.add_page_break()


def viet_bai3(doc):
    v = bai3.lan_truyen_tien(**bai3.GIA_TRI_DAU)
    g = bai3.lan_truyen_nguoc(v)
    doc.add_heading("BÀI 3. ĐỒ THỊ TÍNH TOÁN VÀ LAN TRUYỀN NGƯỢC", level=1)
    doan(doc, "Hàm số J = 3a + 2(4b + 5cd) với a = 2, b = 3, c = 4, d = 5.")

    doc.add_heading("a) Phân rã thành các phép toán cơ bản", level=2)
    bang(
        doc,
        ["Biến trung gian", "Phép toán", "Loại"],
        [["u₁ = 3a", "nhân hằng số", "×"], ["u₂ = 4b", "nhân hằng số", "×"], ["u₃ = c·d", "nhân hai biến", "×"], ["u₄ = 5u₃", "nhân hằng số", "×"], ["u₅ = u₂ + u₄", "cộng", "+"], ["u₆ = 2u₅", "nhân hằng số", "×"], ["J = u₁ + u₆", "cộng", "+"]],
    )

    doc.add_heading("b) Đồ thị tính toán", level=2)
    hinh(doc, bai3.ve_do_thi_bai3(v, g), "Hình 7. Đồ thị tính toán của J = 3a + 2(4b + 5cd): giá trị tiến (xanh) và gradient ngược (đỏ)", rong=16.5)
    gach_dau_dong(doc, "Nút đầu vào (hình chữ nhật xanh): a, b, c, d.")
    gach_dau_dong(doc, "Nút phép toán (hình tròn cam): ×3, ×4, × (c·d), ×5, +, ×2, + tạo ra u₁, …, u₆ và J.")
    gach_dau_dong(doc, "Nút đầu ra (hình chữ nhật xanh lá): J. Cạnh có hướng biểu diễn sự phụ thuộc: a → u₁ → J; b → u₂ → u₅; c, d → u₃ → u₄ → u₅ → u₆ → J.")

    doc.add_heading("c) Lan truyền tiến", level=2)
    bang(
        doc,
        ["Thứ tự", "Nút", "Công thức", "Giá trị"],
        [
            [1, "u₁", "3a = 3·2", f"{v['u1']:g}"],
            [2, "u₂", "4b = 4·3", f"{v['u2']:g}"],
            [3, "u₃", "c·d = 4·5", f"{v['u3']:g}"],
            [4, "u₄", "5u₃ = 5·20", f"{v['u4']:g}"],
            [5, "u₅", "u₂ + u₄ = 12 + 100", f"{v['u5']:g}"],
            [6, "u₆", "2u₅ = 2·112", f"{v['u6']:g}"],
            [7, "J", "u₁ + u₆ = 6 + 224", f"{v['J']:g}"],
        ],
    )

    doc.add_heading("d) Đạo hàm cục bộ tại từng nút", level=2)
    doan(doc, "Quy tắc chung: với phép cộng p = x + y thì ∂p/∂x = ∂p/∂y = 1 (nút cộng “phân phối” gradient nguyên vẹn); với phép nhân p = x·y thì ∂p/∂x = y, ∂p/∂y = x (nút nhân “hoán đổi” giá trị đầu vào); với p = k·x thì ∂p/∂x = k.")
    bang(
        doc,
        ["Nút", "Đạo hàm cục bộ", "Giá trị"],
        [
            ["u₁ = 3a", "∂u₁/∂a = 3", "3"],
            ["u₂ = 4b", "∂u₂/∂b = 4", "4"],
            ["u₃ = c·d", "∂u₃/∂c = d;  ∂u₃/∂d = c", "5;  4"],
            ["u₄ = 5u₃", "∂u₄/∂u₃ = 5", "5"],
            ["u₅ = u₂ + u₄", "∂u₅/∂u₂ = 1;  ∂u₅/∂u₄ = 1", "1;  1"],
            ["u₆ = 2u₅", "∂u₆/∂u₅ = 2", "2"],
            ["J = u₁ + u₆", "∂J/∂u₁ = 1;  ∂J/∂u₆ = 1", "1;  1"],
        ],
    )

    doc.add_heading("e) Lan truyền ngược", level=2)
    bang(
        doc,
        ["Bước", "Gradient", "Quy tắc dây chuyền", "Giá trị"],
        [
            [0, "∂J/∂J", "khởi tạo", "1"],
            [1, "∂J/∂u₁", "∂J/∂J · ∂J/∂u₁ = 1·1", f"{g['u1']:g}"],
            [2, "∂J/∂u₆", "∂J/∂J · ∂J/∂u₆ = 1·1", f"{g['u6']:g}"],
            [3, "∂J/∂u₅", "∂J/∂u₆ · ∂u₆/∂u₅ = 1·2", f"{g['u5']:g}"],
            [4, "∂J/∂u₂", "∂J/∂u₅ · ∂u₅/∂u₂ = 2·1", f"{g['u2']:g}"],
            [5, "∂J/∂u₄", "∂J/∂u₅ · ∂u₅/∂u₄ = 2·1", f"{g['u4']:g}"],
            [6, "∂J/∂u₃", "∂J/∂u₄ · ∂u₄/∂u₃ = 2·5", f"{g['u3']:g}"],
            [7, "∂J/∂a", "∂J/∂u₁ · ∂u₁/∂a = 1·3", f"{g['a']:g}"],
            [8, "∂J/∂b", "∂J/∂u₂ · ∂u₂/∂b = 2·4", f"{g['b']:g}"],
            [9, "∂J/∂c", "∂J/∂u₃ · ∂u₃/∂c = 10·d = 10·5", f"{g['c']:g}"],
            [10, "∂J/∂d", "∂J/∂u₃ · ∂u₃/∂d = 10·c = 10·4", f"{g['d']:g}"],
        ],
    )
    cong_thuc(doc, f"∂J/∂a = {g['a']:g},   ∂J/∂b = {g['b']:g},   ∂J/∂c = {g['c']:g},   ∂J/∂d = {g['d']:g}")

    doc.add_heading("f) Kiểm chứng bằng khai triển trực tiếp", level=2)
    cong_thuc(doc, "J = 3a + 2(4b + 5cd) = 3a + 8b + 10cd")
    cong_thuc(doc, "∂J/∂a = 3,   ∂J/∂b = 8,   ∂J/∂c = 10d = 50,   ∂J/∂d = 10c = 40")
    doan(doc, "Kết quả hoàn toàn trùng khớp với lan truyền ngược trên đồ thị ở câu (e). Kiểm tra giá trị: J = 3·2 + 8·3 + 10·4·5 = 6 + 24 + 200 = 230.")

    doc.add_heading("g) Một bước Gradient Descent với α = 0.01", level=2)
    moi = bai3.cap_nhat_gradient_descent(bai3.GIA_TRI_DAU, g)
    J_moi = bai3.ham_J(**moi)
    bang(doc, ["Tham số", "Giá trị cũ", "Gradient", "Cập nhật", "Giá trị mới"], [[k, f"{bai3.GIA_TRI_DAU[k]:g}", f"{g[k]:g}", f"{bai3.GIA_TRI_DAU[k]:g} − 0.01·{g[k]:g}", f"{moi[k]:.2f}"] for k in ("a", "b", "c", "d")])
    cong_thuc(doc, f"J′ = 3·{moi['a']:.2f} + 8·{moi['b']:.2f} + 10·{moi['c']:.2f}·{moi['d']:.2f} = {3 * moi['a']:.2f} + {8 * moi['b']:.2f} + {10 * moi['c'] * moi['d']:.2f} = {J_moi:.2f}")
    doan(doc, f"Nhận xét: J giảm từ 230 xuống {J_moi:.2f} (giảm {v['J'] - J_moi:.2f}, khoảng {100 * (v['J'] - J_moi) / v['J']:.1f}%). Xấp xỉ bậc nhất dự đoán mức giảm α·‖∇J‖² = 0.01·(3² + 8² + 50² + 40²) = 0.01·4173 = 41.73; mức giảm thực tế 39.73 nhỏ hơn một chút do số hạng bậc hai α²·10·(∂J/∂c)(∂J/∂d) = 0.0001·10·50·40 = 2 của tích cd. Các biến c, d có gradient lớn nhất nên thay đổi nhiều nhất. Vì J là hàm tuyến tính theo a, b nên J không có cực tiểu, lặp tiếp sẽ làm J giảm mãi.")

    doc.add_heading("h) Kiểm chứng bằng PyTorch Autograd", level=2)
    J_pt, g_pt = bai3.tinh_bang_pytorch(bai3.GIA_TRI_DAU)
    khoi_code(doc, inspect.getsource(bai3.tinh_bang_pytorch))
    bang(doc, ["Đại lượng", "Tính thủ công", "PyTorch Autograd", "Khớp"], [["J", f"{v['J']:g}", f"{J_pt:g}", "✓" if J_pt == v["J"] else "✗"]] + [[f"∂J/∂{k}", f"{g[k]:g}", f"{g_pt[k]:g}", "✓" if g_pt[k] == g[k] else "✗"] for k in ("a", "b", "c", "d")])
    doan(doc, "PyTorch xây dựng đồ thị tính toán động khi thực hiện các phép toán trên tensor có requires_grad=True; lời gọi J.backward() duyệt đồ thị theo thứ tự ngược và áp dụng quy tắc dây chuyền giống hệt câu (e), lưu kết quả vào thuộc tính .grad. Kết quả trùng khớp hoàn toàn với cách tính thủ công.")


def phu_luc(doc):
    doc.add_page_break()
    doc.add_heading("PHỤ LỤC. MÃ NGUỒN", level=1)
    doan(doc, f"Toàn bộ mã nguồn nằm trong thư mục BT02_ComputationGraph trên nhánh BT02_ComputationGraph của kho: {REPO}")
    bang(
        doc,
        ["File", "Nội dung"],
        [
            ["do_thi_tinh_toan.py", "Hàm vẽ đồ thị tính toán (forward/backward) dùng chung"],
            ["bai1_logistic_tre_so_sinh.py", "Bài 1: tính toán 2 bước Gradient Descent, dự đoán, vẽ đồ thị"],
            ["bai2_logistic_danh_gia_phim.py", "Bài 2: chuẩn hóa, 2 bước GD, độ đo, ranh giới, huấn luyện 1000 vòng NumPy/PyTorch"],
            ["bai3_lan_truyen_nguoc.py", "Bài 3: lan truyền tiến/ngược thủ công, Gradient Descent, PyTorch Autograd"],
            ["tao_bao_cao.py", "Sinh file báo cáo Word này"],
        ],
    )
    doan(doc, "Cách chạy:")
    khoi_code(doc, "pip install -r requirements.txt\npython bai1_logistic_tre_so_sinh.py\npython bai2_logistic_danh_gia_phim.py\npython bai3_lan_truyen_nguoc.py\npython tao_bao_cao.py")


def main():
    doc = Document()
    cai_dat_tai_lieu(doc)
    them_so_trang(doc)
    trang_bia(doc)
    viet_bai1(doc)
    viet_bai2(doc)
    viet_bai3(doc)
    phu_luc(doc)
    duong_dan = THU_MUC / TEN_FILE
    doc.save(duong_dan)
    print("Đã tạo:", duong_dan)
    thu_muc_nop = THU_MUC.parent.parent
    if thu_muc_nop.exists():
        shutil.copy2(duong_dan, thu_muc_nop / TEN_FILE)
        print("Đã sao chép:", thu_muc_nop / TEN_FILE)


if __name__ == "__main__":
    np.set_printoptions(precision=6, suppress=True)
    main()
